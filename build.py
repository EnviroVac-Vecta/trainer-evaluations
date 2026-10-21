"""Build the trainer evaluations site.

Writes site/ with:
  index.html, forms.json, vendor/      (copied as-is)
  config.json   trainer list (trainers.txt) + where Save & Send goes (EVAL_SUBMIT_URL secret)
  data.json     employee list for the picker + cohort history (history.json, imported once
                from the Trainer Tracker). The employee list comes from the flow in the
                EVAL_DATA_URL secret when set, otherwise from the imported roster.

Local test with a file instead of the flow:
    python build.py "Employee file.xlsx"
"""
import io
import os
import datetime as dt
import re
import sys
import json
import shutil

import requests
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "site")
SECRETS = json.loads(os.environ.get("ALL_SECRETS") or "{}")


def secret(name):
    return SECRETS.get(name) or os.environ.get(name) or ""


def norm(text):
    return "".join(ch for ch in str(text or "").lower() if ch.isalnum())


# HSI system/test logins that aren't real employees.
SYSTEM_ACCOUNT = re.compile(r"donesafe|^test\b|\btest student\b|limited access", re.I)

ROSTER_COLUMNS = {
    "name": ["name", "employeename", "employee", "fullname", "trainee", "membername"],
    "first": ["firstname", "first"],
    "last": ["lastname", "last"],
    "id": ["employeeid", "empid", "id", "loginid", "eeid"],
    "office": ["office", "homeoffice", "location", "team", "teams", "branch"],
    "status": ["status", "employeestatus"],
}


def get_data_file():
    if len(sys.argv) > 1:
        with open(sys.argv[1], "rb") as f:
            return f.read()
    url = secret("EVAL_DATA_URL")
    if not url:
        print("No EVAL_DATA_URL secret yet: employee names will be typed in.")
        return None
    res = requests.post(url, json={}, timeout=180)
    if res.status_code != 200 or not res.content.startswith(b"PK"):
        raise SystemExit(f"The employee-file flow didn't send back an Excel file (status {res.status_code}).")
    return res.content


def read_roster(xlsx_bytes):
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for h, row in enumerate(rows[:15]):
            headers = [norm(c) for c in row]
            found = {}
            for key, names in ROSTER_COLUMNS.items():
                for want in names:
                    if want in headers:
                        found[key] = headers.index(want)
                        break
            if "name" in found or ("first" in found and "last" in found):
                people = {}
                for r in rows[h + 1:]:
                    get = lambda k: (str(r[found[k]]).strip() if k in found and found[k] < len(r) and r[found[k]] is not None else "")
                    if "status" in found and get("status").lower() not in ("active", ""):
                        continue  # HSI member list: only Active people
                    last, emp_id = get("last"), get("id")
                    if "," in last:  # HSI puts the employee ID after the last name: "Chase,F3IH3KEL"
                        last, tail = [x.strip() for x in last.split(",", 1)]
                        if re.fullmatch(r"[A-Z0-9]{6,12}", tail):
                            emp_id = emp_id if emp_id and "@" not in emp_id else tail
                    if "@" in emp_id:
                        emp_id = ""
                    name = get("name") or (get("first") + " " + last).strip()
                    if not name or SYSTEM_ACCOUNT.search(name):
                        continue
                    key = emp_id or name.lower()
                    people.setdefault(key, {"name": name, "id": emp_id, "office": get("office")})
                roster = sorted(people.values(), key=lambda p: p["name"].lower())
                print(f"Employee list: {len(roster)} people from sheet '{ws.title}'.")
                return roster
    raise SystemExit("Couldn't find a Name column (or First/Last Name) in the employee file.")


# ---------------- Cohort Daily Report (HSI Envirovac Cohort Status Report) ----------------
REQ_SUFFIX = re.compile(r"\s*\(\s*\d+\s*(m|mo|yr|y)\s*\)\s*$", re.I)
# When HSI lists the same training twice (someone in two cohort groups), the best status wins:
# a completion date, then still in its training window, then delinquent.
STATUS_RANK = {"complete": 3, "Incomplete": 2, "Delinquent": 1}
COHORT_ORDER = ["Tech I to Tech II", "Tech II to Tech III", "Tech III to Jr. Operator", "Jr. Operator to Sr. Operator",
                "Sr. Operator to Crew Leader", "Crew Leader to Supervisor"]


def get_cohort_file():
    path = os.environ.get("COHORT_FILE")
    if path:
        with open(path, "rb") as f:
            return f.read()
    url = secret("EVAL_COHORT_URL")
    if not url:
        print("No EVAL_COHORT_URL secret yet: cohort tabs use the imported history only.")
        return None
    res = requests.post(url, json={}, timeout=180)
    if res.status_code != 200 or not res.content.startswith(b"PK"):
        raise SystemExit(f"The cohort-report flow didn't send back an Excel file (status {res.status_code}).")
    return res.content


def name_key(n):
    n = str(n or "")
    if "," in n:
        last, first = n.split(",", 1)
        n = first + " " + last
    t = [x for x in re.sub(r"[^a-z ]", " ", n.lower()).split() if len(x) > 1]
    return f"{t[0]} {t[-1]}" if t else ""


def read_cohort_report(xlsx_bytes):
    """-> {key: {name, id, office, cohort, trainings{col: date|Incomplete|Delinquent}}}"""
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows[:15]) if r and "First Name" in [str(v).strip() for v in r if v]), None)
    if hi is None:
        raise SystemExit("Cohort report: couldn't find the header row (First Name, Last Name, ...).")
    h = [str(v).strip() if v else "" for v in rows[hi]]
    col = {c: h.index(c) for c in ("First Name", "Last Name", "Teams", "Groups", "Requirement", "Compliance Status", "Last Completion Date")}
    people = {}
    for r in rows[hi + 1:]:
        if not r or not r[col["First Name"]]:
            continue
        g = lambda c: r[col[c]] if col[c] < len(r) else None
        last, emp_id = str(g("Last Name") or ""), ""
        if "," in last:
            last, emp_id = [x.strip() for x in last.split(",", 1)]
        name = f"{str(g('First Name')).strip()} {last}".strip()
        key = emp_id or name_key(name)
        groups = [re.sub(r"^Envirovac Cohort\s*", "", x.strip()) for x in str(g("Groups") or "").split(",") if x.strip()]
        p = people.setdefault(key, {"name": name, "id": emp_id, "office": str(g("Teams") or ""), "groups": set(), "trainings": {}})
        p["groups"].update(groups)
        req = REQ_SUFFIX.sub("", str(g("Requirement") or "")).strip()
        status = str(g("Compliance Status") or "").strip().lower()
        done = g("Last Completion Date")
        if status == "current" and done:
            d = done if isinstance(done, (dt.date, dt.datetime)) else None
            val, rank = (d.strftime("%m/%d/%y") if d else str(done)[:10]), STATUS_RANK["complete"]
        elif status == "delinquent":
            val, rank = "Delinquent", STATUS_RANK["Delinquent"]
        else:
            val, rank = "Incomplete", STATUS_RANK["Incomplete"]
        old = p["trainings"].get(req)
        if old is None or rank > STATUS_RANK.get("complete" if re.match(r"\d", old) else old, 0):
            p["trainings"][req] = val
    for p in people.values():
        gs = sorted(p.pop("groups"), key=lambda c: COHORT_ORDER.index(c) if c in COHORT_ORDER else -1)
        p["cohorts"] = gs
        p["cohort"] = gs[-1] if gs else ""   # highest level, unless the history already knows which one
    print(f"Cohort report: {len(people)} people, {sum(len(p['trainings']) for p in people.values())} requirements.")
    return people


def merge_cohort_report(history, report):
    """Update the cohort history with today's HSI report. People stay in the history after they
    drop off the report, so past cohorts are kept."""
    by_id = {p["id"]: p for p in history["people"] if p.get("id")}
    by_name = {name_key(p["name"]): p for p in history["people"]}
    cols = history["trainings"]
    added, updated = [], 0
    for key, r in report.items():
        p = by_id.get(r["id"]) if r["id"] else None
        p = p or by_name.get(name_key(r["name"]))
        if p is None:
            p = {"name": r["name"], "id": r["id"], "office": r["office"], "jobTitle": "", "target": "",
                 "cohortDates": "", "start": None, "end": None, "period": "", "employed": True,
                 "notRecommended": False, "cohort": r["cohort"], "trainings": {}, "inHSIReport": True}
            history["people"].append(p)
            added.append(r["name"])
        for t in r["trainings"]:
            if t not in cols:
                cols.append(t)
        p["trainings"] = {t: r["trainings"].get(t, "N/A") for t in cols}
        p["inHSIReport"] = True
        p["employed"] = True
        p["lastSeen"] = dt.date.today().isoformat()
        if r["cohort"] and p.get("cohort") not in r["cohorts"]:
            p["cohort"] = r["cohort"]
        if r["office"]:
            p["office"] = r["office"]
        if not p.get("id"):
            p["id"] = r["id"]
        updated += 1
    # Rebuild "which trainings each cohort requires" from the people HSI currently reports.
    req = {}
    for p in history["people"]:
        if p.get("inHSIReport") and p.get("cohort"):
            req.setdefault(p["cohort"], set()).update(t for t, v in p["trainings"].items() if v and v.upper() != "N/A")
    history["requiredByCohort"] = {c: [t for t in cols if t in v] for c, v in req.items()}
    history["cohortReportDate"] = dt.date.today().isoformat()
    print(f"Merged cohort report: {updated} people updated, {len(added)} new: {added[:10]}")
    print(f"::notice title=Cohort report::{updated} people updated from HSI, {len(added)} new to a cohort"
          + (f": {', '.join(added[:15])}" if added else ""))
    return history


def send_trainer_report(history):
    """Trainer copy (colored Excel) to the "Trainer Report - Save file" flow -> Project Polarity, same name every day.
    Never published on the public site."""
    url = secret("TRAINER_REPORT_URL")
    if not url:
        print("No TRAINER_REPORT_URL secret yet: trainer copy of the Training Cycle report not sent.")
        return
    import base64
    from tools.trainer_report import build
    name = "Training Cycle Report.xlsx"
    body = {"fileName": name, "fileContent": base64.b64encode(build(history)).decode("ascii")}
    try:
        r = requests.post(url, json=body, timeout=120)
        if r.status_code in (200, 202):
            print(f"::notice title=Trainer report::Sent '{name}' to Project Polarity.")
        else:
            print(f"::warning title=Trainer report::The flow answered {r.status_code}: {r.text[:300]}")
    except Exception as e:
        print(f"::warning title=Trainer report::Couldn't send it ({e}).")


def main():
    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(OUT)
    for name in ("index.html", "forms.json"):
        shutil.copy(os.path.join(HERE, name), OUT)
    shutil.copytree(os.path.join(HERE, "vendor"), os.path.join(OUT, "vendor"))

    trainers = []
    path = os.path.join(HERE, "trainers.txt")
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            trainers = [l.strip() for l in f if l.strip() and not l.lstrip().startswith("#")]
    submit = secret("EVAL_SUBMIT_URL")
    if not submit:
        print("::warning title=Save & Send not connected::Add the EVAL_SUBMIT_URL secret; until then trainers can only download the PDF.")
    with open(os.path.join(OUT, "config.json"), "w") as f:
        json.dump({"trainers": trainers, "submitUrl": submit, "logUrl": secret("EVAL_LOG_URL")}, f)

    with open(os.path.join(HERE, "history.json")) as f:
        history = json.load(f)
    cohort = get_cohort_file()
    if cohort:
        history = merge_cohort_report(history, read_cohort_report(cohort))
        # Save the updated history so people are remembered after they leave the report.
        with open(os.path.join(HERE, "history.json"), "w") as f:
            json.dump(history, f, indent=1)
    send_trainer_report(history)
    data = get_data_file()
    roster = read_roster(data) if data else history.pop("roster", [])
    history.pop("roster", None)
    with open(os.path.join(OUT, "data.json"), "w") as f:
        json.dump({"roster": roster, "history": history}, f, separators=(",", ":"))
    with open(os.path.join(OUT, "robots.txt"), "w") as f:
        f.write("User-agent: *\nDisallow: /\n")
    print(f"::notice title=Trainer evaluations::{len(trainers)} trainers, Save & Send {'on' if submit else 'off'}, "
          f"{len(roster)} employees, {len(history['people'])} cohort members")


if __name__ == "__main__":
    main()

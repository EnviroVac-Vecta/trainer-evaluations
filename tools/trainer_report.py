"""Trainer copy of the Training Cycle report (Excel), sent daily to the Project Polarity folder.

One row per person, one column per training. Built from the same history the trainer
website uses (HSI Cohort Daily Report merged in), so past training cycles are included.

Colors:
  training cells  light green = current (completed, shows the date)
                  light yellow = needed (not done yet, not delinquent)
                  light red    = delinquent
  employee name   light blue = in a current training cycle
                  light gray = past training cycle (ended or no longer employed)
"""
import io
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN, YELLOW, RED = "E2F0D9", "FFF2CC", "F8D7DA"
BLUE, GRAY = "DDEBF7", "E7E6E6"
HEAD = "203864"
DATE_RE = re.compile(r"^\d{1,2}/\d{1,2}/\d{2,4}$")


def _fill(c):
    return PatternFill("solid", fgColor=c)


def is_current(p, today):
    """Same rule as the website's Current Training Cycles tab."""
    if not p.get("employed", True):
        return False
    return p["end"] >= today if p.get("end") else bool(p.get("inHSIReport"))


def period(p):
    if p.get("period"):
        return p["period"]
    f = lambda iso: datetime.strptime(iso, "%Y-%m-%d").strftime("%b %Y") if iso else ""
    return f"{f(p.get('start'))} - {f(p.get('end'))}" if p.get("start") else ""


def build(history, today=None):
    """Return the workbook as bytes."""
    today = today or date.today().isoformat()
    people = history["people"]
    trainings = [t for t in history["trainings"]
                 if any((p.get("trainings") or {}).get(t, "N/A").upper() != "N/A" for p in people)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Training Cycles"
    head = ["Employee", "Employee ID", "Office", "Job Title", "Training Cycle", "Cycle Period", "Current / Past",
            "Intake Form", "Complete", "Needed", "Delinquent"] + trainings
    ws.append(head)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _fill(HEAD)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = box
    ws.row_dimensions[1].height = 60

    rows = sorted(people, key=lambda p: (not is_current(p, today), p.get("cohort") or "", p.get("name", "")))
    first_t = len(head) - len(trainings) + 1
    for p in rows:
        cur = is_current(p, today)
        tr = p.get("trainings") or {}
        cells, fills = [], []
        done = need = late = 0
        for t in trainings:
            v = str(tr.get(t, "N/A") or "N/A").strip()
            if DATE_RE.match(v):
                cells.append(v); fills.append(GREEN); done += 1
            elif v.lower().startswith("delinq"):
                cells.append("Delinquent"); fills.append(RED); late += 1
            elif v.upper() == "N/A":
                cells.append(""); fills.append(None)
            else:
                cells.append("Needed"); fills.append(YELLOW); need += 1
        intake = "Complete" if (p.get("intakeDate") or p.get("intake")) else "Incomplete"
        ws.append([p.get("name", ""), p.get("id", ""), p.get("office", ""), p.get("jobTitle", ""), p.get("cohort", ""),
                   period(p), "Current" if cur else "Past", intake, done, need, late] + cells)
        r = ws.max_row
        ws.cell(r, 1).fill = _fill(BLUE if cur else GRAY)
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 8).fill = _fill(GREEN if intake == "Complete" else YELLOW)
        for i, f in enumerate(fills):
            if f:
                ws.cell(r, first_t + i).fill = _fill(f)
        for c in ws[r]:
            c.border = box
            if c.column >= 8:
                c.alignment = Alignment(horizontal="center")

    widths = [28, 13, 18, 30, 24, 20, 11, 11, 10, 9, 11] + [14] * len(trainings)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{ws.max_row}"

    key = wb.create_sheet("Key")
    key.append(["Training Cycle report", f"from the HSI Cohort Daily Report, {datetime.strptime(today, '%Y-%m-%d').strftime('%B %d, %Y')}"])
    key["A1"].font = Font(bold=True, size=13)
    key.append([])
    for label, color, text in (
        ("Employee name", BLUE, "In a current training cycle"),
        ("Employee name", GRAY, "Past training cycle (ended, or no longer employed)"),
        ("Date", GREEN, "Current: completed on that date"),
        ("Needed", YELLOW, "Required, not completed yet (not delinquent)"),
        ("Delinquent", RED, "Required and past due"),
        ("(blank)", None, "Not required for this person's training cycle"),
        ("Intake Form", None, "Complete or Incomplete"),
    ):
        key.append([label, text])
        if color:
            key.cell(key.max_row, 1).fill = _fill(color)
    key.append([])
    key.append(["Tip", "Use the filter arrows in row 1 (for example Current / Past, Office, or a training column)."])
    key.column_dimensions["A"].width = 18
    key.column_dimensions["B"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    import json, sys
    h = json.load(open(sys.argv[1] if len(sys.argv) > 1 else "history.json"))
    out = sys.argv[2] if len(sys.argv) > 2 else "Training Cycle Report.xlsx"
    open(out, "wb").write(build(h))
    print("wrote", out)

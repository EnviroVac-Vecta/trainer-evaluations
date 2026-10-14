"""One-time import of cohort history from the Trainer Tracker into history.json.

    python tools/import_tracker.py "Trainer_Tracker.xlsx"

Keeps training fields only (no emails, comments, or termination details).
"""
import re, sys, json, datetime as dt
from openpyxl import load_workbook

COHORT_FROM_TARGET = {
    "tech 2": "Tech I to Tech II", "t2": "Tech I to Tech II",
    "tech 3": "Tech II to Tech III", "t3": "Tech II to Tech III",
    "jr. operator": "Tech III to Jr. Operator", "jo": "Tech III to Jr. Operator",
    "sr. operator": "Jr. Operator to Sr. Operator", "so": "Jr. Operator to Sr. Operator",
    "crew leader": "Sr. Operator to Crew Leader", "cl": "Sr. Operator to Crew Leader",
    "supervisor": "Crew Leader to Supervisor", "s": "Crew Leader to Supervisor",
}

def first_last(n):
    n = str(n or "").strip()
    if "," in n:
        last, first = n.split(",", 1)
        return (first.strip() + " " + last.strip()).strip()
    return n

def key(n):
    t = [x for x in re.sub(r"[^a-z ]", " ", first_last(n).lower()).split() if len(x) > 1]
    return f"{t[0]} {t[-1]}" if t else ""

def parse_period(text):
    """'4 Month 9.1.2026 to 12.31.2026' -> ('2026-09-01', '2026-12-31')"""
    m = re.findall(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", str(text or ""))
    if len(m) < 2:
        return None, None
    out = []
    for i, (mo, d, y) in enumerate(m[:2]):
        mo, d, y = int(mo), int(d), int(y)
        for day in (d, 30, 28):
            try:
                out.append(dt.date(y, mo, day).isoformat()); break
            except ValueError:
                continue
    return out[0], out[1]

def cell(v):
    if v is None: return ""
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%m/%d/%y")
    return str(v).strip()

def main(path):
    wb = load_workbook(path, data_only=True)
    people = {}
    def rows(sheet):
        r = list(wb[sheet].iter_rows(values_only=True)); h = [str(x or "").strip() for x in r[0]]
        return h, [x for x in r[1:] if any(v not in (None, "") for v in x)]
    for sheet, employed in (("Active Roster", True), ("Terminated - Cohort", False), ("Terminated - No Cohort", False)):
        h, data = rows(sheet)
        col = lambda name: next((i for i, x in enumerate(h) if x.lower() == name.lower()), next((i for i, x in enumerate(h) if x.lower().startswith(name.lower())), None))
        ci = {k: col(k) for k in ("Employee", "Employee ID", "Team", "Job Title", "Target to Evaluate", "Cohort Dates", "Promtion Recommendation", "Ready to be Promoted")}
        for r in data:
            g = lambda k: (r[ci[k]] if ci[k] is not None and ci[k] < len(r) else None)
            if not g("Employee"): continue
            start, end = parse_period(g("Cohort Dates"))
            rec = str(g("Promtion Recommendation") or "").strip().lower()
            people.setdefault(key(g("Employee")), {
                "name": first_last(g("Employee")), "id": cell(g("Employee ID")), "office": cell(g("Team")),
                "jobTitle": cell(g("Job Title")), "target": cell(g("Target to Evaluate")),
                "cohortDates": cell(g("Cohort Dates")), "start": start, "end": end,
                "employed": employed, "notRecommended": rec == "not recommended",
                "cohort": "", "trainings": {}, "inHSIReport": False,
            })
    h, data = rows("Training Matrix")
    trainings = h[4:]
    for r in data:
        if not r[2]: continue
        k = key(r[2]); p = people.get(k)
        if not p:
            p = people.setdefault(k, {"name": str(r[2]).strip(), "id": "", "office": cell(r[3]), "jobTitle": "", "target": "",
                                      "cohortDates": "", "start": None, "end": None, "employed": True, "notRecommended": False,
                                      "cohort": "", "trainings": {}, "inHSIReport": False})
        p["cohort"] = cell(r[0]); p["period"] = cell(r[1]); p["inHSIReport"] = True
        if not p["office"]: p["office"] = cell(r[3])
        p["trainings"] = {t: cell(v) for t, v in zip(trainings, r[4:]) if t}
    for p in people.values():
        if not p["cohort"] and p["start"]:
            p["cohort"] = COHORT_FROM_TARGET.get(p["target"].lower(), "")
    cohort_people = [p for p in people.values() if p["start"]]
    # Which trainings each cohort requires (anything not N/A for that cohort in the matrix).
    required = {}
    for p in cohort_people:
        if p["inHSIReport"]:
            req = required.setdefault(p["cohort"], set())
            req.update(t for t, v in p["trainings"].items() if v.upper() != "N/A")
    roster = sorted(({"name": p["name"], "id": p["id"], "office": p["office"]} for p in people.values() if p["employed"]),
                    key=lambda r: r["name"].lower())
    out = {
        "roster": roster,
        "importedFrom": path.split("/")[-1], "importedAt": dt.date.today().isoformat(),
        "trainings": trainings,
        "requiredByCohort": {c: sorted(v, key=trainings.index) for c, v in required.items()},
        "people": sorted(cohort_people, key=lambda p: (p["start"] or "", p["name"])),
    }
    with open("history.json", "w") as f:
        json.dump(out, f, indent=1)
    missing = [p["name"] for p in cohort_people if not p["inHSIReport"]]
    print(f"{len(cohort_people)} people with a cohort; {sum(1 for p in cohort_people if not p['employed'])} no longer employed;"
          f" not in HSI report: {missing}; not recommended: {[p['name'] for p in cohort_people if p['notRecommended']]}")
    print("cohorts:", {c: len(v) for c, v in out["requiredByCohort"].items()})

if __name__ == "__main__":
    main(sys.argv[1])
